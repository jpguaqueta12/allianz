"""
Carga puntual del backlog desde ../backlog.xlsx hacia Azure SQL.

La carga es directa a base de datos:
- Backlog Fabrica -> dbo.backlog_fabrica
- Backlog Mejora Continua -> dbo.backlog_mejora_continua

El PI de Fabrica comparte el PI activo de MEJORA_CONTINUA, segun la logica de
la aplicacion. El identificador funcional del backlog en este Excel viene en
Labels (IBLCDM-*), no en Issue key, y se usa como ticket_key para respetar la
convencion ya existente en la base.
"""
from __future__ import annotations

import argparse
import asyncio
import json
import re
import sys
import unicodedata
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
BACKEND_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(BACKEND_ROOT))
DEFAULT_EXCEL = ROOT / "backlog.xlsx"

from app.db.connection import close_pool, get_pool, init_pool

SHEETS = {
    "Backlog Fabrica": "dbo.backlog_fabrica",
    "Backlog Mejora Continua": "dbo.backlog_mejora_continua",
}

HEADER_ALIASES = {
    "issue_type": ["issue type", "tipo"],
    "ticket_key": ["labels", "label", "ibl", "ticket", "issue key"],
    "summary": ["summary", "resumen", "descripcion", "descripción", "titulo", "título"],
    "status": ["status", "estado"],
    "resolution": ["resolution", "resolucion", "resolución"],
    "responsable_java": ["asignado java", "responsable java"],
    "responsable_cobol": ["asignado cobol", "responsable cobol"],
    "responsable_qa": ["asignado calidad", "asigna calidad", "responsable calidad", "responsable qa"],
    "horas_desarrollo_java": ["horas java"],
    "horas_desarrollo_cobol": ["horas cobol"],
    "horas_analisis_qa": ["horas calidad", "horas qa"],
    "fecha_finalizacion_inicial": ["fecha compromiso", "fecha comprometida"],
    "fecha_escalado": ["fecha escalamiento", "fecha escalado"],
    "fecha_reinicio": ["fecha de desescalamiento", "fecha desescalamiento", "fecha reinicio"],
    "fecha_finalizacion": ["nueva fecha de entrega", "fecha finalizacion", "fecha finalización"],
    "fecha_entrega": ["fecha entrega", "fecha de entrega"],
    "etc": ["etc"],
}

INSERTABLE_COLUMNS = [
    "issue_type",
    "ticket_key",
    "status",
    "resolution",
    "summary",
    "pi_id",
    "responsable_java",
    "responsable_cobol",
    "responsable_qa",
    "horas_desarrollo_java",
    "horas_desarrollo_cobol",
    "horas_analisis_qa",
    "fecha_finalizacion_inicial",
    "fecha_finalizacion",
    "fecha_escalado",
    "fecha_reinicio",
    "fecha_entrega",
    "etc",
    "escalados",
]

TEXT_LIMITS = {
    "issue_type": 100,
    "ticket_key": 80,
    "status": 100,
    "resolution": 120,
    "responsable_java": 300,
    "responsable_cobol": 300,
    "responsable_qa": 300,
}

RESPONSABLE_ALIASES = {
    "alvaro alfonso lasso": "Alvaro Alfonso Lasso Lopez",
    "andres felipe novoa": "Andres Felipe Novoa Garcia",
    "andres novoa": "Andres Felipe Novoa Garcia",
    "angie lizeth cordoba": "Angie Lizeth Cordoba Lesmes",
    "camilo lobo guerrero": "Camilo Lobo Guerrero Nova",
    "carlos andres pavajeau": "Carlos Andres Pavajeau Max",
    "carlos villadiego": "Carlos Villadiego",
    "deivis david sanchez": "Deivis David Sanchez Mestra",
    "diego alejandro rodriguez": "Diego Alejandro Rodriguez Martinez",
    "dilan camilo marinez": "Dilan Camilo Martinez Zapata",
    "dilan camilo martinez": "Dilan Camilo Martinez Zapata",
    "dylan camilo martinez zapata": "Dilan Camilo Martinez Zapata",
    "erick steven alegria": "Erik Steven Alegria Mina",
    "erik steven alegria": "Erik Steven Alegria Mina",
    "fernanda pardo": "Laura Fernanda Pardo",
    "fredy fernando patino": "Fredy Fernando Pati\u00f1o Rave",
    "heidy vanesa sanchez": "Heidy Vanessa Sanchez Pulido",
    "jeison stiven rojas": "Jeison Stiven Rojas Montoya",
    "jaisson steven rojas": "Jeison Stiven Rojas Montoya",
    "jeisson andres cutiva": "Jeisson Andres Cutiva Cardenas",
    "jhon carlos colorado": "Jhon Carlos Colorado Angulo",
    "johan david garzon": "Johan David Garzon Uricoechea",
    "john robledo": "John Jairo Robledo Quintero",
    "jorge enrique castillo": "Jorge Enrique Castillo Gonzalez",
    "juan carlos villareal": "Juan Carlos Villarreal Carrera",
    "juan carlos villarreal": "Juan Carlos Villarreal Carrera",
    "juan david caceres": "Juan David Caceres Aponte",
    "kevin alejandro correa": "Kevin Alejandro Correa Hurtado",
    "laura fernanda pardo": "Laura Fernanda Pardo",
    "luisa fernanda pardo": "Laura Fernanda Pardo",
    "maria fernanda alvarado": "Maria Fernanda Alvarado",
    "maryerin hernandez": "Maryerin Hernandez",
    "nicolas cardenas": "Nicolas Cardenas Rodriguez",
    "nicolas menaca": "Nicolas Andres Menaca Trujillo",
    "rafael alvarado": "Rafael Alvarado",
    "rafael enrique alvarado": "Rafael Alvarado",
    "roger armando lozada": "Roger Armando Lozada Ortiz",
    "sandra lorena marinez": "Sandra Lorena Martinez Merchan",
    "sandra lorena martinez": "Sandra Lorena Martinez Merchan",
    "santiago nicolas brisnez": "Santiago Nicolas Bri\u00f1ez Garcia",
    "santiago nicolas brinez": "Santiago Nicolas Bri\u00f1ez Garcia",
    "sergio panche": "Sergio Alejandro Panche",
}

PERSONA_TECH = {
    "Alvaro Alfonso Lasso Lopez": "JAVA",
    "Andres Felipe Novoa Garcia": "JAVA",
    "Angie Lizeth Cordoba Lesmes": "COBOL",
    "Camilo Lobo Guerrero Nova": "JAVA",
    "Carlos Andres Pavajeau Max": "JAVA",
    "Carlos Villadiego": "CALIDAD",
    "Deivis David Sanchez Mestra": "JAVA",
    "Diego Alejandro Rodriguez Martinez": "JAVA",
    "Dilan Camilo Martinez Zapata": "JAVA",
    "Erik Steven Alegria Mina": "JAVA",
    "Fredy Fernando Pati\u00f1o Rave": "COBOL",
    "Heidy Vanessa Sanchez Pulido": "COBOL",
    "Jeison Stiven Rojas Montoya": "JAVA",
    "Jeisson Andres Cutiva Cardenas": "COBOL",
    "Jhon Carlos Colorado Angulo": "JAVA",
    "Johan David Garzon Uricoechea": "JAVA",
    "John Jairo Robledo Quintero": "JAVA",
    "Jorge Enrique Castillo Gonzalez": "JAVA",
    "Juan Carlos Villarreal Carrera": "COBOL",
    "Juan David Caceres Aponte": "COBOL",
    "Kevin Alejandro Correa Hurtado": "COBOL",
    "Laura Fernanda Pardo": "CALIDAD",
    "Maria Fernanda Alvarado": "COBOL",
    "Maryerin Hernandez": "CALIDAD",
    "Nicolas Andres Menaca Trujillo": "JAVA",
    "Nicolas Cardenas Rodriguez": "JAVA",
    "Rafael Alvarado": "CALIDAD",
    "Roger Armando Lozada Ortiz": "COBOL",
    "Sandra Lorena Martinez Merchan": "COBOL",
    "Santiago Nicolas Bri\u00f1ez Garcia": "JAVA",
    "Sergio Alejandro Panche": "COBOL",
}


def normalize_label(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value.strip().lower())
    without_accents = "".join(c for c in normalized if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", without_accents).strip()


def fix_mojibake(value: str | None) -> str | None:
    if value is None:
        return None
    cleaned = str(value).strip()
    if not cleaned:
        return None
    if cleaned.upper() in {"N/A", "NA", "NULL", "NONE", "PENDIENTE ASIGNACION"}:
        return None
    if "Ã" in cleaned or "â" in cleaned:
        try:
            return cleaned.encode("latin1").decode("utf-8")
        except UnicodeError:
            return cleaned
    return cleaned


def limit_text(field: str, value: str | None) -> str | None:
    limit = TEXT_LIMITS.get(field)
    if value is not None and limit and len(value) > limit:
        return value[:limit]
    return value


def normalize_responsables(value: Any, expected_tech: str) -> str | None:
    cleaned = fix_mojibake(value)
    if not cleaned:
        return None
    parts = re.split(r"\s*(?:\||;|,|/|\by\b)\s*", cleaned, flags=re.IGNORECASE)
    result: list[str] = []
    for raw_part in parts:
        part = re.sub(r"\s+-\s*validar\s*$", "", raw_part.strip(), flags=re.IGNORECASE)
        key = normalize_label(part)
        if key in {"verificar", "validar", "pendiente asignacion", "na", "n a"}:
            continue
        match = None
        for full_name, full_tech in PERSONA_TECH.items():
            if full_tech == expected_tech and normalize_label(full_name) == key:
                match = full_name
                break
        if not match:
            match = RESPONSABLE_ALIASES.get(key)
        if not match:
            continue
        if PERSONA_TECH.get(match) != expected_tech:
            continue
        if match not in result:
            result.append(match)
    return " | ".join(result) if result else None


def expected_tech_for_profile(perfil: str | None) -> str | None:
    normalized = (perfil or "").strip().lower()
    if normalized == "java":
        return "JAVA"
    if normalized == "cobol":
        return "COBOL"
    if normalized in {"calidad", "qa"}:
        return "CALIDAD"
    return None


def normalize_planificacion_extra(extra: Any) -> tuple[str | None, bool]:
    if not extra:
        return None, False
    try:
        data = json.loads(extra) if isinstance(extra, str) else extra
    except json.JSONDecodeError:
        return extra, False
    if not isinstance(data, dict):
        return extra, False
    items = data.get("planificacion_items")
    if not isinstance(items, list):
        return extra, False

    changed = False
    normalized_items = []
    for item in items:
        if not isinstance(item, dict):
            normalized_items.append(item)
            continue
        expected_tech = expected_tech_for_profile(item.get("perfil"))
        normalized_name = normalize_responsables(item.get("responsable"), expected_tech) if expected_tech else item.get("responsable")
        if normalized_name != item.get("responsable"):
            changed = True
        normalized_items.append({**item, "responsable": normalized_name})
    data["planificacion_items"] = normalized_items
    return json.dumps(data, ensure_ascii=False, default=str), changed


def to_decimal(value: Any) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, (int, float, Decimal)):
        return Decimal(str(value))
    cleaned = str(value).strip().replace(",", ".")
    if cleaned.upper() in {"N/A", "NA"}:
        return None
    try:
        return Decimal(cleaned)
    except (InvalidOperation, ValueError):
        return None


def to_int(value: Any) -> int | None:
    dec = to_decimal(value)
    return int(dec) if dec is not None else None


def to_date(value: Any) -> date | None:
    if value is None or str(value).strip() == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%d/%m/%Y %H:%M", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            pass
    return None


def resolve_headers(headers: list[str]) -> dict[str, int]:
    normalized_headers = [normalize_label(h) for h in headers]
    resolved: dict[str, int] = {}
    for field, aliases in HEADER_ALIASES.items():
        for alias in aliases:
            wanted = normalize_label(alias)
            if wanted in normalized_headers:
                resolved[field] = normalized_headers.index(wanted)
                break
    return resolved


def cell(row: tuple[Any, ...], index: int | None) -> Any:
    if index is None or index >= len(row):
        return None
    return row[index]


def total_hours(row: dict[str, Any]) -> Decimal:
    total = Decimal("0")
    for field in ("horas_desarrollo_java", "horas_desarrollo_cobol", "horas_analisis_qa"):
        total += row.get(field) or Decimal("0")
    return total


def row_score(row: dict[str, Any]) -> tuple[int, Decimal, int, int]:
    issue_type = (row.get("issue_type") or "").strip().lower()
    return (
        0 if issue_type == "estimacion" else 1,
        total_hours(row),
        1 if row.get("fecha_finalizacion") or row.get("fecha_finalizacion_inicial") else 0,
        int(row["_source_row"]),
    )


def read_sheet(path: Path, sheet_name: str) -> tuple[list[dict[str, Any]], int, int]:
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name]
    raw_headers = next(ws.iter_rows(values_only=True))
    headers = [str(h).strip() if h is not None and str(h).strip() else f"Col{i + 1}" for i, h in enumerate(raw_headers)]
    indexes = resolve_headers(headers)

    parsed: list[dict[str, Any]] = []
    skipped_no_key = 0
    for row_number, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not any(v is not None and str(v).strip() for v in row):
            continue

        ticket_key = fix_mojibake(cell(row, indexes.get("ticket_key")))
        if not ticket_key:
            skipped_no_key += 1
            continue

        fecha_fin_inicial = to_date(cell(row, indexes.get("fecha_finalizacion_inicial")))
        fecha_fin = to_date(cell(row, indexes.get("fecha_finalizacion"))) or fecha_fin_inicial
        fecha_escalado = to_date(cell(row, indexes.get("fecha_escalado")))
        fecha_reinicio = to_date(cell(row, indexes.get("fecha_reinicio")))
        escalados = None
        if fecha_escalado:
            escalados = json.dumps([{
                "fecha_escalado": fecha_escalado.isoformat(),
                "fecha_reinicio": fecha_reinicio.isoformat() if fecha_reinicio else None,
            }], ensure_ascii=False)

        summary = fix_mojibake(cell(row, indexes.get("summary"))) or ticket_key
        parsed.append({
            "_source_row": row_number,
            "issue_type": limit_text("issue_type", fix_mojibake(cell(row, indexes.get("issue_type")))),
            "ticket_key": limit_text("ticket_key", ticket_key.upper()),
            "summary": summary,
            "status": limit_text("status", fix_mojibake(cell(row, indexes.get("status")))),
            "resolution": limit_text("resolution", fix_mojibake(cell(row, indexes.get("resolution")))),
            "responsable_java": limit_text("responsable_java", normalize_responsables(cell(row, indexes.get("responsable_java")), "JAVA")),
            "responsable_cobol": limit_text("responsable_cobol", normalize_responsables(cell(row, indexes.get("responsable_cobol")), "COBOL")),
            "responsable_qa": limit_text("responsable_qa", normalize_responsables(cell(row, indexes.get("responsable_qa")), "CALIDAD")),
            "horas_desarrollo_java": to_decimal(cell(row, indexes.get("horas_desarrollo_java"))),
            "horas_desarrollo_cobol": to_decimal(cell(row, indexes.get("horas_desarrollo_cobol"))),
            "horas_analisis_qa": to_decimal(cell(row, indexes.get("horas_analisis_qa"))),
            "fecha_finalizacion_inicial": fecha_fin_inicial,
            "fecha_finalizacion": fecha_fin,
            "fecha_escalado": fecha_escalado,
            "fecha_reinicio": fecha_reinicio,
            "fecha_entrega": to_date(cell(row, indexes.get("fecha_entrega"))),
            "etc": to_int(cell(row, indexes.get("etc"))) or 0,
            "escalados": escalados,
        })
    wb.close()

    by_key: dict[str, dict[str, Any]] = {}
    duplicate_rows = 0
    for row in parsed:
        key = row["ticket_key"]
        if key in by_key:
            duplicate_rows += 1
            if row_score(row) > row_score(by_key[key]):
                by_key[key] = row
        else:
            by_key[key] = row

    return list(by_key.values()), skipped_no_key, duplicate_rows


async def table_columns(pool: Any, table: str) -> set[str]:
    schema, name = table.split(".", 1)
    rows = await pool.fetch(
        """
        SELECT COLUMN_NAME
        FROM INFORMATION_SCHEMA.COLUMNS
        WHERE TABLE_SCHEMA=$1 AND TABLE_NAME=$2
        """,
        schema,
        name,
    )
    return {r["COLUMN_NAME"] for r in rows}


async def upsert_row(conn: Any, table: str, columns: list[str], row: dict[str, Any], pi_id: int, dry_run: bool) -> str:
    values = {col: row.get(col) for col in columns if col != "pi_id"}
    values["pi_id"] = pi_id
    exists = await conn.fetchval(
        f"SELECT COUNT(*) FROM {table} WHERE ticket_key=$1 AND pi_id=$2",
        row["ticket_key"],
        pi_id,
    )
    if dry_run:
        return "update" if exists else "insert"

    if exists:
        update_columns = [c for c in columns if c not in {"ticket_key", "pi_id"}]
        assignments = ", ".join(f"{col}=${idx + 3}" for idx, col in enumerate(update_columns))
        await conn.execute(
            f"UPDATE {table} SET {assignments} WHERE ticket_key=$1 AND pi_id=$2",
            row["ticket_key"],
            pi_id,
            *(values[col] for col in update_columns),
        )
        return "update"

    insert_cols = ", ".join(columns)
    placeholders = ", ".join(f"${idx + 1}" for idx in range(len(columns)))
    await conn.execute(
        f"INSERT INTO {table} ({insert_cols}) VALUES ({placeholders})",
        *(values[col] for col in columns),
    )
    return "insert"


async def normalize_table_planificacion(conn: Any, table: str, pi_id: int, dry_run: bool) -> int:
    rows = await conn.fetch(
        f"""
        SELECT id, extra, responsable_java, responsable_cobol, responsable_qa
        FROM {table}
        WHERE pi_id=$1
        """,
        pi_id,
    )
    changed = 0
    for row in rows:
        next_java = normalize_responsables(row["responsable_java"], "JAVA")
        next_cobol = normalize_responsables(row["responsable_cobol"], "COBOL")
        next_qa = normalize_responsables(row["responsable_qa"], "CALIDAD")
        next_extra, extra_changed = normalize_planificacion_extra(row["extra"])
        fields_changed = (
            next_java != row["responsable_java"]
            or next_cobol != row["responsable_cobol"]
            or next_qa != row["responsable_qa"]
            or extra_changed
        )
        if not fields_changed:
            continue
        changed += 1
        if dry_run:
            continue
        await conn.execute(
            f"""
            UPDATE {table}
            SET responsable_java=$1,
                responsable_cobol=$2,
                responsable_qa=$3,
                extra=$4
            WHERE id=$5
            """,
            next_java,
            next_cobol,
            next_qa,
            next_extra,
            row["id"],
        )
    return changed


async def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--file", type=Path, default=DEFAULT_EXCEL)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    if not args.file.exists():
        raise SystemExit(f"No existe el archivo: {args.file}")

    await init_pool()
    pool = get_pool()
    pi_id = await pool.fetchval(
        "SELECT id FROM dbo.pi WHERE activo=1 AND modulo='MEJORA_CONTINUA' ORDER BY fecha_inicio DESC, id DESC LIMIT 1"
    )
    if not pi_id:
        raise SystemExit("No hay PI activo de MEJORA_CONTINUA. Fabrica comparte ese PI.")

    summary: dict[str, dict[str, int]] = {}
    async with pool.acquire() as conn:
        for sheet_name, table in SHEETS.items():
            db_columns = await table_columns(conn, table)
            columns = [col for col in INSERTABLE_COLUMNS if col in db_columns]
            rows, skipped_no_key, duplicate_rows = read_sheet(args.file, sheet_name)

            inserted = updated = 0
            for row in rows:
                action = await upsert_row(conn, table, columns, row, int(pi_id), args.dry_run)
                if action == "insert":
                    inserted += 1
                else:
                    updated += 1
            normalized = await normalize_table_planificacion(conn, table, int(pi_id), args.dry_run)

            summary[sheet_name] = {
                "filas_unicas": len(rows),
                "insertadas": inserted,
                "actualizadas": updated,
                "omitidas_sin_ticket_key": skipped_no_key,
                "duplicados_consolidados": duplicate_rows,
                "normalizadas_planificacion": normalized,
            }

    await close_pool()
    print(json.dumps({"pi_id": pi_id, "dry_run": args.dry_run, "summary": summary}, ensure_ascii=False, indent=2))
    return 0


if __name__ == "__main__":
    raise SystemExit(asyncio.run(main()))
