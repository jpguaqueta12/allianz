from __future__ import annotations
import io
import json
import re
import unicodedata
from datetime import datetime
from decimal import Decimal, InvalidOperation
from fastapi import APIRouter, HTTPException, UploadFile, File
import structlog

logger = structlog.get_logger()
router = APIRouter()

MAX_FILE_SIZE = 20 * 1024 * 1024  # 20 MB

# Mapeo de nombres de columna del Excel → campo interno
# Soporta variaciones de Jira (inglés/español, con espacios o guiones)
_COL_ALIASES: dict[str, list[str]] = {
    "ticket_key":            ["issue key", "key", "issuekey", "id", "ticket", "issue id", "ibl", "id ibl", "ticket ibl", "codigo ibl", "código ibl"],
    "summary":               ["summary", "resumen", "título", "titulo", "title", "nombre", "descripcion", "descripción", "detalle", "asunto"],
    "issue_type":            ["issue type", "issuetype", "tipo", "type"],
    "status":                ["status", "estado"],
    "priority":              ["priority", "prioridad"],
    "resolution":            ["resolution", "resolución", "resolucion"],
    "assignee":              ["assignee", "asignado", "assigned to", "asignado a"],
    "reporter":              ["reporter", "reportador", "reported by"],
    "created":               ["created", "creado", "fecha creación", "fecha creacion", "created date"],
    "updated":               ["updated", "actualizado", "updated date", "fecha actualización"],
    "epic_link":             ["epic link", "epic_link", "epica", "épica", "epic name", "enlace épica", "epic"],
    "story_points":          ["story points", "story point estimate", "puntos historia", "sp", "estimado", "story point"],
    "sprint":                ["sprint", "sprints", "sprint name"],
    "labels":                ["labels", "etiquetas", "label"],
    "components":            ["components", "componentes", "component"],
    "fix_version":           ["fix version", "fix version/s", "versión", "version", "versiones"],
    "project":               ["project", "proyecto", "project name", "project key"],
    "include_release_notes": ["include to release notes", "include release notes", "release notes", "incluir en notas"],
    "assigned_team":         ["assigned team", "equipo asignado", "team", "equipo"],
}


def _normalize_label(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value.strip().lower())
    without_accents = "".join(c for c in normalized if not unicodedata.combining(c))
    return re.sub(r"[^a-z0-9]+", " ", without_accents).strip()


def _resolve_columns(headers: list[str]) -> dict[str, int | None]:
    mapping: dict[str, int | None] = {f: None for f in _COL_ALIASES}
    normalized_aliases = {
        field: {_normalize_label(alias) for alias in aliases}
        for field, aliases in _COL_ALIASES.items()
    }
    for i, h in enumerate(headers):
        hl = _normalize_label(h)
        for field, aliases in _COL_ALIASES.items():
            if mapping[field] is None and hl in normalized_aliases[field]:
                mapping[field] = i
    return mapping


def _column_mapping_quality(mapping: dict[str, int | None]) -> int:
    important = ("ticket_key", "summary", "status", "epic_link", "assignee", "story_points")
    return sum(1 for field in important if mapping.get(field) is not None)


def _needs_ai_column_mapping(mapping: dict[str, int | None]) -> bool:
    return (
        mapping.get("ticket_key") is None
        or mapping.get("summary") is None
        or _column_mapping_quality(mapping) < 3
    )


def _cell_to_str(value) -> str:
    return str(value).strip() if value is not None else ""


def _extract_json_object(text: str) -> dict | None:
    cleaned = text.strip()
    cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned, flags=re.IGNORECASE)
    cleaned = re.sub(r"\s*```$", "", cleaned)
    match = re.search(r"\{.*\}", cleaned, flags=re.DOTALL)
    if match:
        cleaned = match.group(0)
    try:
        parsed = json.loads(cleaned)
        return parsed if isinstance(parsed, dict) else None
    except json.JSONDecodeError:
        return None


def _resolve_ai_header_index(value, headers: list[str]) -> int | None:
    if value is None:
        return None
    if isinstance(value, int):
        if 0 <= value < len(headers):
            return value
        if 1 <= value <= len(headers):
            return value - 1
        return None

    wanted = _normalize_label(str(value))
    if not wanted or wanted in {"null", "none", "ninguno"}:
        return None
    normalized_headers = [_normalize_label(h) for h in headers]
    for i, header in enumerate(normalized_headers):
        if wanted == header:
            return i
    for i, header in enumerate(normalized_headers):
        if wanted and (wanted in header or header in wanted):
            return i
    return None


async def _resolve_columns_with_ai(
    headers: list[str],
    data_rows: list,
    heuristic_mapping: dict[str, int | None],
) -> tuple[dict[str, int | None], dict]:
    if not _needs_ai_column_mapping(heuristic_mapping):
        return heuristic_mapping, {"used_ai": False, "confidence": None, "reason": None}

    sample_rows = []
    for row in data_rows[:8]:
        sample_rows.append({
            headers[i]: _cell_to_str(row[i])[:120] if i < len(row) else ""
            for i in range(len(headers))
            if headers[i]
        })

    fields = {
        "ticket_key": "Identificador unico del ticket o IBL, por ejemplo IBLCDM-22569, MD-123 o FAST-45.",
        "summary": "Resumen, titulo, asunto o descripcion corta del requerimiento.",
        "issue_type": "Tipo de issue, historia, bug, tarea, requerimiento.",
        "status": "Estado del ticket.",
        "priority": "Prioridad.",
        "resolution": "Resolucion.",
        "assignee": "Responsable o asignado.",
        "reporter": "Reportador o solicitante.",
        "created": "Fecha de creacion.",
        "updated": "Fecha de actualizacion.",
        "epic_link": "Epica, iniciativa, proyecto, frente o referencia usada para clasificar MD/Fabrica.",
        "story_points": "Puntos de historia, esfuerzo o estimacion numerica.",
        "sprint": "Sprint.",
        "labels": "Etiquetas.",
        "components": "Componentes.",
        "fix_version": "Version o release.",
        "project": "Proyecto.",
        "include_release_notes": "Indicador de notas de release.",
        "assigned_team": "Equipo asignado.",
    }
    prompt = (
        "Eres un asistente que mapea columnas de Excel a campos internos de Jira/backlog.\n"
        "Devuelve solo JSON valido. No inventes columnas: usa exactamente un encabezado recibido o null.\n"
        "Campos internos y significado:\n"
        f"{json.dumps(fields, ensure_ascii=False)}\n\n"
        "Formato requerido:\n"
        "{\"mapping\":{\"ticket_key\":\"encabezado o null\", ...},\"confidence\":0.0,\"reason\":\"breve\"}\n\n"
        f"Encabezados: {json.dumps(headers, ensure_ascii=False)}\n"
        f"Muestra de filas: {json.dumps(sample_rows, ensure_ascii=False, default=str)}"
    )

    try:
        from app.services.llm_service import get_llm

        response = await get_llm().ainvoke(prompt)
        content = getattr(response, "content", "")
        if isinstance(content, list):
            content = " ".join(str(part.get("text", part)) if isinstance(part, dict) else str(part) for part in content)
        parsed = _extract_json_object(str(content))
        ai_mapping = parsed.get("mapping") if parsed else None
        if not isinstance(ai_mapping, dict):
            raise ValueError("La IA no devolvio un mapping valido")

        mapping = dict(heuristic_mapping)
        for field in _COL_ALIASES:
            if mapping.get(field) is not None:
                continue
            idx = _resolve_ai_header_index(ai_mapping.get(field), headers)
            if idx is not None:
                mapping[field] = idx

        return mapping, {
            "used_ai": True,
            "confidence": parsed.get("confidence"),
            "reason": parsed.get("reason"),
        }
    except Exception as exc:
        logger.warning("backlog_ai_column_mapping_failed", error=str(exc))
        return heuristic_mapping, {
            "used_ai": False,
            "confidence": None,
            "reason": f"No se pudo usar IA para mapear columnas: {exc}",
        }


async def _resolve_columns_for_rows(all_rows: list) -> tuple[list[str], dict[str, int | None], dict]:
    raw_headers = [str(h).strip() if h is not None else f"Col{i+1}"
                   for i, h in enumerate(all_rows[0])]
    heuristic = _resolve_columns(raw_headers)
    mapping, ai_meta = await _resolve_columns_with_ai(raw_headers, all_rows[1:], heuristic)
    mapped_headers = {
        field: raw_headers[idx]
        for field, idx in mapping.items()
        if idx is not None and idx < len(raw_headers)
    }
    meta = {
        **ai_meta,
        "mapped_headers": mapped_headers,
        "mapped_fields_count": len(mapped_headers),
    }
    return raw_headers, mapping, meta


def _extract_project_key(epic_link: str) -> str:
    """'MD-123' → 'MD',  'MIPYME-45' → 'MIPYME',  'MD' → 'MD'"""
    m = re.match(r'^([A-Z][A-Z0-9_]*)', epic_link.strip().upper())
    return m.group(1) if m else epic_link.strip().upper()


async def _build_modulo_lookup(pool) -> dict[str, str]:
    """Devuelve {identi_proyecto: modulo} para todos los proyectos activos."""
    rows = await pool.fetch("SELECT identi, modulo FROM proyectos WHERE activo = TRUE")
    return {r["identi"]: r["modulo"] for r in rows}


def _classify(epic_link: str | None, lookup: dict[str, str]) -> str:
    """
    Extrae el prefijo del Epic Link y lo busca en los identificadores de proyecto.
    Si no hay match devuelve 'FABRICA' como default.
    """
    if not epic_link or not epic_link.strip():
        return "FABRICA"
    key = _extract_project_key(epic_link)
    return lookup.get(key, "FABRICA")


def _safe_decimal(value: str) -> Decimal | None:
    try:
        return Decimal(str(value).replace(",", "."))
    except (InvalidOperation, ValueError):
        return None


def _safe_datetime(value: str) -> datetime | None:
    for fmt in ("%d/%b/%y %I:%M %p", "%Y-%m-%dT%H:%M:%S", "%Y-%m-%d %H:%M:%S",
                "%d/%m/%Y %H:%M", "%m/%d/%Y %H:%M", "%Y-%m-%d"):
        try:
            return datetime.strptime(value.strip(), fmt)
        except ValueError:
            continue
    return None


def _parse_rows(
    all_rows: list,
    modulo_lookup: dict[str, str],
    col_map: dict[str, int | None] | None = None,
) -> tuple[list[str], list[dict]]:
    """Convierte las filas crudas del Excel en dicts clasificados."""
    raw_headers = [str(h).strip() if h is not None else f"Col{i+1}"
                   for i, h in enumerate(all_rows[0])]
    col_map = col_map or _resolve_columns(raw_headers)
    known_indices = {v for v in col_map.values() if v is not None}

    rows: list[dict] = []
    for source_row, raw in enumerate(all_rows[1:], start=2):
        cells = [str(v).strip() if v is not None else "" for v in raw]
        if not any(cells):           # fila completamente vacía → saltar
            continue

        def get(field: str) -> str | None:
            idx = col_map.get(field)
            return cells[idx] if idx is not None and idx < len(cells) and cells[idx] else None

        epic_val  = get("epic_link")
        modulo    = _classify(epic_val, modulo_lookup)
        sp_raw    = get("story_points")
        cr_raw    = get("created")
        up_raw    = get("updated")

        extra = {raw_headers[i]: cells[i]
                 for i in range(len(cells))
                 if i not in known_indices and cells[i]}

        rows.append({
            "created":               _safe_datetime(cr_raw) if cr_raw else None,
            "issue_type":            get("issue_type"),
            "ticket_key":            get("ticket_key"),
            "project":               get("project"),
            "status":                get("status"),
            "include_release_notes": get("include_release_notes"),
            "resolution":            get("resolution"),
            "summary":               get("summary") or "(sin resumen)",
            "assigned_team":         get("assigned_team"),
            "assignee":              get("assignee"),
            "reporter":              get("reporter"),
            "epic_link":             epic_val,
            "priority":              get("priority"),
            "updated":               _safe_datetime(up_raw) if up_raw else None,
            "story_points":          _safe_decimal(sp_raw) if sp_raw else None,
            "sprint":                get("sprint"),
            "labels":                get("labels"),
            "components":            get("components"),
            "fix_version":           get("fix_version"),
            "modulo":                modulo,
            "extra":                 extra or None,
            "_source_row":           source_row,
        })

    return raw_headers, rows


def _normalize_ticket_key(value: str | None) -> str | None:
    if not value:
        return None
    normalized = re.sub(r"\s+", "", value.strip().upper())
    return normalized or None


def _find_file_duplicates(rows: list[dict]) -> list[dict]:
    seen: dict[str, dict] = {}
    for row in rows:
        key = _normalize_ticket_key(row.get("ticket_key"))
        if not key:
            continue
        entry = seen.setdefault(key, {"ticket_key": key, "veces": 0, "filas": []})
        entry["veces"] += 1
        entry["filas"].append(row.get("_source_row"))
    return [entry for entry in seen.values() if entry["veces"] > 1]


async def _find_existing_duplicates(pool, pi_id: int, keys: list[str]) -> list[dict]:
    normalized_keys = sorted({k for k in (_normalize_ticket_key(key) for key in keys) if k})
    if not normalized_keys:
        return []

    duplicates: list[dict] = []
    tables = [
        ("MEJORA_CONTINUA", "backlog_mejora_continua"),
        ("FABRICA", "backlog_fabrica"),
    ]
    for chunk_start in range(0, len(normalized_keys), 400):
        chunk = normalized_keys[chunk_start:chunk_start + 400]
        placeholders = ",".join(f"${i + 2}" for i in range(len(chunk)))
        for modulo, table in tables:
            rows = await pool.fetch(
                f"""
                SELECT ticket_key
                FROM {table}
                WHERE pi_id=$1 AND UPPER(REPLACE(ticket_key, ' ', '')) IN ({placeholders})
                """,
                pi_id,
                *chunk,
            )
            duplicates.extend({
                "ticket_key": _normalize_ticket_key(r["ticket_key"]) or r["ticket_key"],
                "modulo": modulo,
            } for r in rows)
    return duplicates


def _find_header_row(rows: list) -> int:
    """
    Busca en las primeras 10 filas cuál es la fila de encabezados.
    Devuelve su índice. Usa la que tenga más coincidencias con las columnas conocidas.
    """
    all_aliases = {_normalize_label(alias) for aliases in _COL_ALIASES.values() for alias in aliases}
    best_idx, best_score = 0, 0
    for i, row in enumerate(rows[:10]):
        score = sum(
            1 for cell in row
            if cell is not None and _normalize_label(str(cell)) in all_aliases
        )
        if score > best_score:
            best_score, best_idx = score, i
    return best_idx


def _ensure_ticket_key_mapping(col_map: dict[str, int | None], mapping_meta: dict) -> None:
    if col_map.get("ticket_key") is not None:
        return
    detail = (
        "No se pudo identificar la columna del IBL/ticket. Renombra la columna a IBL, Ticket o Issue Key, "
        "o revisa la configuración de Azure OpenAI para que la IA pueda mapear el archivo."
    )
    reason = mapping_meta.get("reason")
    if reason:
        detail = f"{detail} Detalle: {reason}"
    raise HTTPException(status_code=422, detail=detail)


def _read_excel_rows(content: bytes) -> list:
    """
    Carga el Excel, detecta automáticamente la fila de encabezados
    y devuelve [header_row, data_row1, data_row2, ...].
    """
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    target = wb.worksheets[0]
    for ws in wb.worksheets:
        if "ticket" in ws.title.lower():
            target = ws
            break
    all_rows = list(target.iter_rows(values_only=True))
    wb.close()

    if not all_rows:
        raise ValueError("El archivo está vacío")

    header_idx = _find_header_row(all_rows)
    rows = all_rows[header_idx:]          # header + datos

    if len(rows) < 2:
        raise ValueError("No se encontraron filas de datos después del encabezado")
    return rows


# ── INSERT helpers ─────────────────────────────────────────────────────────────

_INSERT_MC = """
    INSERT INTO backlog_mejora_continua
        (created, issue_type, ticket_key, project, status, include_release_notes,
         resolution, summary, assigned_team, assignee, reporter, epic_link,
         priority, updated, story_points, sprint, labels, components, fix_version,
         pi_id, extra)
    VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11,$12,$13,$14,$15,$16,$17,$18,$19,$20,$21)
"""

_INSERT_FAB = """
    INSERT INTO backlog_fabrica
        (created, issue_type, ticket_key, project, status, include_release_notes,
         resolution, summary, assigned_team, assignee, reporter, epic_link,
         priority, updated, story_points, sprint, labels, components, fix_version,
         pi_id, extra)
    VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9,$10,$11,$12,$13,$14,$15,$16,$17,$18,$19,$20,$21)
"""


def _row_params(r: dict, pi_id: int) -> tuple:
    return (
        r["created"], r["issue_type"], r["ticket_key"], r["project"],
        r["status"], r["include_release_notes"], r["resolution"], r["summary"],
        r["assigned_team"], r["assignee"], r["reporter"], r["epic_link"],
        r["priority"], r["updated"], r["story_points"],
        r["sprint"], r["labels"], r["components"], r["fix_version"],
        pi_id,
        json.dumps(r["extra"], ensure_ascii=False, default=str) if r["extra"] else None,
    )


# ── Endpoints ──────────────────────────────────────────────────────────────────

@router.post("/upload/preview-excel")
async def preview_excel(file: UploadFile = File(...)):
    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="El archivo supera el límite de 20 MB")
    ext = (file.filename or "").lower()
    if not any(ext.endswith(e) for e in (".xlsx", ".xls", ".xlsm")):
        raise HTTPException(status_code=422, detail="Solo se admiten archivos Excel (.xlsx, .xls, .xlsm)")

    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        sheets = []
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            all_rows = list(ws.iter_rows(values_only=True))
            if not all_rows:
                sheets.append({"nombre": sheet_name, "columnas": [], "filas": [], "total_filas": 0})
                continue
            columns = [str(h) if h is not None else f"Col {i+1}" for i, h in enumerate(all_rows[0])]
            filas = [[str(v) if v is not None else "" for v in row] for row in all_rows[1:501]]
            sheets.append({"nombre": sheet_name, "columnas": columns,
                            "filas": filas, "total_filas": len(all_rows) - 1})
        wb.close()

        tiene_epic_link = any(
            any("epic link" in (c or "").lower() for c in s["columnas"]) for s in sheets
        )
        return {"filename": file.filename, "hojas": sheets, "tiene_epic_link": tiene_epic_link}
    except HTTPException:
        raise
    except Exception as exc:
        raise HTTPException(status_code=422, detail=f"No se pudo leer el archivo: {exc}")


@router.post("/upload/debug-excel")
async def debug_excel(file: UploadFile = File(...)):
    """Diagnóstico: muestra encabezados reales, columnas detectadas y muestra de epic_links."""
    content = await file.read()
    raw_rows = _read_excel_rows(content)
    raw_headers, col_map, mapping_meta = await _resolve_columns_for_rows(raw_rows)

    from app.db.connection import get_pool
    modulo_lookup = await _build_modulo_lookup(get_pool())

    # Muestra primeros 20 valores únicos de la columna epic_link
    epic_idx = col_map.get("epic_link")
    epic_samples: list[str] = []
    if epic_idx is not None:
        seen: set[str] = set()
        for row in raw_rows[1:]:
            cells = [str(v).strip() if v is not None else "" for v in row]
            val = cells[epic_idx] if epic_idx < len(cells) else ""
            if val and val not in seen:
                seen.add(val)
                epic_samples.append(val)
            if len(epic_samples) >= 20:
                break

    return {
        "encabezados_originales": raw_headers,
        "columnas_mapeadas": {k: v for k, v in col_map.items() if v is not None},
        "columnas_detectadas": mapping_meta["mapped_headers"],
        "mapeo_con_ia": mapping_meta["used_ai"],
        "confianza_ia": mapping_meta["confidence"],
        "epic_link_col_index": epic_idx,
        "epic_link_muestras": epic_samples,
        "proyectos_en_bd": modulo_lookup,
        "total_filas": len(raw_rows) - 1,
    }


@router.post("/upload/analizar-backlog")
async def analizar_backlog(file: UploadFile = File(...)):
    """
    Clasifica los tickets según Epic Link vs. identificadores de proyectos en la BD.
    Devuelve muestra 5+5 sin guardar nada.
    """
    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="El archivo supera el límite de 20 MB")

    from app.db.connection import get_pool
    pool = get_pool()

    try:
        raw_rows = _read_excel_rows(content)
    except ValueError as e:
        raise HTTPException(status_code=422, detail=str(e))

    modulo_lookup = await _build_modulo_lookup(pool)
    _, col_map, mapping_meta = await _resolve_columns_for_rows(raw_rows)
    _ensure_ticket_key_mapping(col_map, mapping_meta)
    _, rows = _parse_rows(raw_rows, modulo_lookup, col_map)

    mc_rows  = [r for r in rows if r["modulo"] == "MEJORA_CONTINUA"]
    fab_rows = [r for r in rows if r["modulo"] == "FABRICA"]
    duplicados_archivo = _find_file_duplicates(rows)
    pi_mc = await pool.fetchval(
        "SELECT id FROM pi WHERE activo=TRUE AND modulo='MEJORA_CONTINUA' LIMIT 1"
    )
    duplicados_bd = await _find_existing_duplicates(
        pool,
        pi_mc,
        [r["ticket_key"] for r in rows if r.get("ticket_key")],
    ) if pi_mc else []

    # Serializar decimals y datetimes para JSON
    def safe(r: dict) -> dict:
        return {k: (str(v) if isinstance(v, (Decimal, datetime)) else v)
                for k, v in r.items() if k not in {"extra", "_source_row"}}

    return {
        "total":          len(rows),
        "total_mc":       len(mc_rows),
        "total_fabrica":  len(fab_rows),
        "muestra_mc":     [safe(r) for r in mc_rows[:5]],
        "muestra_fabrica":[safe(r) for r in fab_rows[:5]],
        "proyectos_detectados": sorted({
            _extract_project_key(r["epic_link"])
            for r in rows if r["epic_link"]
        }),
        "columnas_detectadas": mapping_meta["mapped_headers"],
        "mapeo_con_ia": mapping_meta["used_ai"],
        "confianza_ia": mapping_meta["confidence"],
        "nota_mapeo": mapping_meta["reason"],
        "duplicados_archivo": duplicados_archivo,
        "duplicados_bd": duplicados_bd,
        "total_duplicados": len(duplicados_archivo) + len(duplicados_bd),
    }


@router.post("/upload/importar-backlog")
async def importar_backlog(file: UploadFile = File(...)):
    """
    Importa fila por fila al backlog correspondiente según Epic Link.
    - MD → backlog_mejora_continua
    - Resto → backlog_fabrica
    """
    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="El archivo supera el límite de 20 MB")

    from app.db.connection import get_pool
    pool = get_pool()

    try:
        raw_rows = _read_excel_rows(content)
    except ValueError as e:
        raise HTTPException(status_code=422, detail=str(e))

    pi_mc = await pool.fetchval(
        "SELECT id FROM pi WHERE activo=TRUE AND modulo='MEJORA_CONTINUA' LIMIT 1"
    )
    if not pi_mc:
        raise HTTPException(
            status_code=400,
            detail="No hay un PI activo configurado. Ve a Configuración → PI, crea un PI y actívalo antes de importar.",
        )
    # FABRICA comparte PI con MEJORA_CONTINUA
    pi_fab = pi_mc

    modulo_lookup = await _build_modulo_lookup(pool)
    _, col_map, mapping_meta = await _resolve_columns_for_rows(raw_rows)
    _ensure_ticket_key_mapping(col_map, mapping_meta)
    _, rows = _parse_rows(raw_rows, modulo_lookup, col_map)

    ins_mc = ins_fab = skipped = 0
    seen_in_file: set[str] = set()

    async with pool.acquire() as conn:
        for r in rows:
            normalized_key = _normalize_ticket_key(r.get("ticket_key"))
            if normalized_key and normalized_key in seen_in_file:
                skipped += 1
                continue
            if normalized_key:
                seen_in_file.add(normalized_key)

            modulo = r["modulo"]
            if modulo == "MEJORA_CONTINUA":
                sql, table, pi_id = _INSERT_MC, "backlog_mejora_continua", pi_mc
            else:
                sql, table, pi_id = _INSERT_FAB, "backlog_fabrica", pi_fab

            try:
                if r.get("ticket_key"):
                    exists_same_table = await conn.fetchval(
                        f"SELECT COUNT(*) FROM {table} WHERE UPPER(REPLACE(ticket_key, ' ', ''))=$1 AND pi_id=$2",
                        normalized_key, pi_id,
                    )
                    exists_other_table = await conn.fetchval(
                        """
                        SELECT COUNT(*)
                        FROM backlog_mejora_continua
                        WHERE UPPER(REPLACE(ticket_key, ' ', ''))=$1 AND pi_id=$2
                        """,
                        normalized_key, pi_id,
                    ) if table == "backlog_fabrica" else await conn.fetchval(
                        """
                        SELECT COUNT(*)
                        FROM backlog_fabrica
                        WHERE UPPER(REPLACE(ticket_key, ' ', ''))=$1 AND pi_id=$2
                        """,
                        normalized_key, pi_id,
                    )
                    if exists_same_table or exists_other_table:
                        skipped += 1
                        continue
                res = await conn.execute(sql, *_row_params(r, pi_id))
                inserted_count = int(res.split()[-1]) if res else 0
                if inserted_count > 0:
                    if modulo == "MEJORA_CONTINUA":
                        ins_mc += 1
                    else:
                        ins_fab += 1
                else:
                    skipped += 1
            except Exception as exc:
                logger.warning("backlog_row_skip", ticket=r.get("ticket_key"), error=str(exc))
                skipped += 1

    return {
        "insertados_mc":      ins_mc,
        "insertados_fabrica": ins_fab,
        "duplicados_omitidos": skipped,
        "total_procesados":   len(rows),
    }


# ── Incidentes ─────────────────────────────────────────────────────────────────

_INC_COL_ALIASES: dict[str, list[str]] = {
    "numero":          ["numero", "número", "number", "incident", "incidente", "id"],
    "equipo":          ["equipo", "team", "assigned team", "grupo", "grupo asignado"],
    "fecha_escalado":  ["fecha escalado sn", "fecha escalado", "escalado", "fecha inicio", "opened", "created"],
    "estado_sn":       ["estado sn", "estado", "status", "state"],
    "jira":            ["jira", "jira ticket", "ticket jira", "issue", "key"],
    "comentario":      ["comentario", "comment", "comments", "notas", "notes", "descripcion", "description"],
    "fecha_respuesta": ["fecha respuesta", "respuesta", "resolved", "fecha cierre", "closed", "updated"],
    "dias":            ["dias", "días", "days", "sla", "tiempo", "duration"],
}


def _resolve_inc_columns(headers: list[str]) -> dict[str, int | None]:
    mapping: dict[str, int | None] = {f: None for f in _INC_COL_ALIASES}
    for i, h in enumerate(headers):
        hl = h.strip().lower()
        for field, aliases in _INC_COL_ALIASES.items():
            if mapping[field] is None and hl in aliases:
                mapping[field] = i
    return mapping


def _find_inc_header_row(rows: list) -> int:
    all_aliases = {alias for aliases in _INC_COL_ALIASES.values() for alias in aliases}
    best_idx, best_score = 0, 0
    for i, row in enumerate(rows[:10]):
        score = sum(
            1 for cell in row
            if cell is not None and str(cell).strip().lower() in all_aliases
        )
        if score > best_score:
            best_score, best_idx = score, i
    return best_idx


def _read_inc_rows(content: bytes) -> list:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    all_rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not all_rows:
        raise ValueError("El archivo está vacío")
    header_idx = _find_inc_header_row(all_rows)
    rows = all_rows[header_idx:]
    if len(rows) < 2:
        raise ValueError("No se encontraron filas de datos después del encabezado")
    return rows


def _safe_int(value: str) -> int | None:
    try:
        return int(str(value).replace(",", ".").split(".")[0])
    except (ValueError, TypeError):
        return None


def _parse_inc_rows(all_rows: list) -> tuple[list[str], list[dict]]:
    raw_headers = [str(h).strip() if h is not None else f"Col{i+1}"
                   for i, h in enumerate(all_rows[0])]
    col_map = _resolve_inc_columns(raw_headers)

    rows: list[dict] = []
    for raw in all_rows[1:]:
        cells = [str(v).strip() if v is not None else "" for v in raw]
        if not any(cells):
            continue

        def get(field: str) -> str | None:
            idx = col_map.get(field)
            return cells[idx] if idx is not None and idx < len(cells) and cells[idx] else None

        fe_raw = get("fecha_escalado")
        fr_raw = get("fecha_respuesta")
        rows.append({
            "numero":         get("numero"),
            "equipo":         get("equipo"),
            "fecha_escalado": _safe_datetime(fe_raw) if fe_raw else None,
            "estado_sn":      get("estado_sn"),
            "jira":           get("jira"),
            "comentario":     get("comentario"),
            "fecha_respuesta":_safe_datetime(fr_raw) if fr_raw else None,
            "dias":           _safe_int(get("dias")) if get("dias") else None,
        })
    return raw_headers, rows


_INSERT_INC = """
    INSERT INTO backlog_incidentes
        (numero, equipo, fecha_escalado, estado_sn, jira, comentario, fecha_respuesta, dias, pi_id)
    VALUES ($1,$2,$3,$4,$5,$6,$7,$8,$9)
"""


@router.post("/upload/analizar-incidentes")
async def analizar_incidentes(file: UploadFile = File(...)):
    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="El archivo supera el límite de 20 MB")

    try:
        raw_rows = _read_inc_rows(content)
    except ValueError as e:
        raise HTTPException(status_code=422, detail=str(e))

    _, rows = _parse_inc_rows(raw_rows)

    def safe(r: dict) -> dict:
        return {k: (str(v) if isinstance(v, datetime) else v) for k, v in r.items()}

    return {
        "total":   len(rows),
        "muestra": [safe(r) for r in rows[:10]],
    }


@router.post("/upload/importar-incidentes")
async def importar_incidentes(file: UploadFile = File(...)):
    content = await file.read()
    if len(content) > MAX_FILE_SIZE:
        raise HTTPException(status_code=413, detail="El archivo supera el límite de 20 MB")

    from app.db.connection import get_pool
    pool = get_pool()

    try:
        raw_rows = _read_inc_rows(content)
    except ValueError as e:
        raise HTTPException(status_code=422, detail=str(e))

    _, rows = _parse_inc_rows(raw_rows)

    pi_id = await pool.fetchval(
        "SELECT id FROM pi WHERE activo=TRUE AND modulo='INCIDENTES' LIMIT 1"
    )
    if not pi_id:
        raise HTTPException(
            status_code=400,
            detail="No hay un PI activo para INCIDENTES. Ve a Configuración → PI, crea un PI de tipo INCIDENTES y actívalo.",
        )

    inserted = skipped = 0
    async with pool.acquire() as conn:
        for r in rows:
            try:
                if r.get("numero"):
                    exists = await conn.fetchval(
                        "SELECT COUNT(*) FROM backlog_incidentes WHERE numero=$1 AND pi_id=$2",
                        r["numero"], pi_id,
                    )
                    if exists:
                        skipped += 1
                        continue
                res = await conn.execute(
                    _INSERT_INC,
                    r["numero"], r["equipo"], r["fecha_escalado"], r["estado_sn"],
                    r["jira"], r["comentario"], r["fecha_respuesta"], r["dias"], pi_id,
                )
                inserted_count = int(res.split()[-1]) if res else 0
                if inserted_count > 0:
                    inserted += 1
                else:
                    skipped += 1
            except Exception as exc:
                logger.warning("inc_row_skip", numero=r.get("numero"), error=str(exc))
                skipped += 1

    return {
        "insertados":         inserted,
        "duplicados_omitidos": skipped,
        "total_procesados":   len(rows),
    }
