from __future__ import annotations
import io
import csv
from pathlib import Path
import structlog

logger = structlog.get_logger()


def extract_text(content: bytes, filename: str) -> dict:
    """
    Extracts plain text from the uploaded file.
    Returns {"text": str, "meta": dict}
    """
    ext = Path(filename).suffix.lower()

    try:
        if ext == ".pdf":
            return _from_pdf(content, filename)
        if ext in (".docx", ".doc"):
            return _from_docx(content, filename)
        if ext in (".xlsx", ".xls", ".xlsm"):
            return _from_excel(content, filename)
        if ext == ".csv":
            return _from_csv(content, filename)
        if ext in (".txt", ".md", ".log", ".json", ".xml", ".html", ".htm"):
            return _from_text(content, filename)
        # Unknown — try plain text as fallback
        return _from_text(content, filename)
    except Exception as exc:
        logger.error("file_extraction_error", filename=filename, error=str(exc))
        raise ValueError(f"No se pudo extraer texto de '{filename}': {exc}")


# ── PDF ───────────────────────────────────────────────────────────────────────

def _from_pdf(content: bytes, filename: str) -> dict:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(content))
    pages_text: list[str] = []
    for page in reader.pages:
        t = page.extract_text()
        if t:
            pages_text.append(t.strip())
    text = "\n\n".join(pages_text)
    return {
        "text": text,
        "meta": {"tipo": "PDF", "paginas": len(reader.pages), "filename": filename},
    }


# ── Word ──────────────────────────────────────────────────────────────────────

def _from_docx(content: bytes, filename: str) -> dict:
    import docx
    doc = docx.Document(io.BytesIO(content))
    parts: list[str] = []

    for para in doc.paragraphs:
        t = para.text.strip()
        if t:
            parts.append(t)

    # Also extract tables
    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n".join(parts)
    return {
        "text": text,
        "meta": {"tipo": "Word", "parrafos": len(doc.paragraphs), "filename": filename},
    }


# ── Excel ─────────────────────────────────────────────────────────────────────

def _from_excel(content: bytes, filename: str) -> dict:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    parts: list[str] = []
    total_rows = 0

    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        sheet_lines: list[str] = [f"=== Hoja: {sheet_name} ==="]
        for row in ws.iter_rows(values_only=True):
            cells = [str(c) for c in row if c is not None and str(c).strip()]
            if cells:
                sheet_lines.append(" | ".join(cells))
                total_rows += 1
            if total_rows > 2000:
                sheet_lines.append("... (truncado: más de 2000 filas)")
                break
        if len(sheet_lines) > 1:
            parts.append("\n".join(sheet_lines))

    wb.close()
    text = "\n\n".join(parts)
    return {
        "text": text,
        "meta": {"tipo": "Excel", "hojas": len(wb.sheetnames), "filas": total_rows, "filename": filename},
    }


# ── CSV ───────────────────────────────────────────────────────────────────────

def _from_csv(content: bytes, filename: str) -> dict:
    text_io = io.StringIO(content.decode("utf-8", errors="replace"))
    reader = csv.reader(text_io)
    lines: list[str] = []
    for i, row in enumerate(reader):
        if i > 2000:
            lines.append("... (truncado)")
            break
        row_clean = [c.strip() for c in row if c.strip()]
        if row_clean:
            lines.append(" | ".join(row_clean))
    text = "\n".join(lines)
    return {
        "text": text,
        "meta": {"tipo": "CSV", "filas": len(lines), "filename": filename},
    }


# ── Plain text ────────────────────────────────────────────────────────────────

def _from_text(content: bytes, filename: str) -> dict:
    text = content.decode("utf-8", errors="replace")
    ext = Path(filename).suffix.lower()
    return {
        "text": text,
        "meta": {"tipo": ext.lstrip(".").upper() or "TXT", "chars": len(text), "filename": filename},
    }
